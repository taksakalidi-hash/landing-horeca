import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.comments import Comment

OUT = "/home/user/landing-horeca/holiday-planner/Holiday_Hosting_Planner.xlsx"

GREEN = "1F4E3D"
CRAN = "9B2335"
CREAM = "FBF7EF"
INPUT = "FFF2CC"
GOLD = "C9A227"
GREY = "8C8C8C"
LINE = "E3DCCB"

F = "Arial"
thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

def font(**kw):
    kw.setdefault("name", F)
    kw.setdefault("size", 10)
    return Font(**kw)

def fill(c):
    return PatternFill("solid", fgColor=c)

wb = Workbook()

def new_sheet(title, tab, widths):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["A"].width = 3
    for r in range(1, 120):
        for c in range(1, 16):
            ws.cell(r, c).fill = fill(CREAM)
            ws.cell(r, c).font = font()
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws

def title(ws, text, sub, last_col):
    ws.row_dimensions[2].height = 34
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last_col)
    c = ws.cell(2, 2, text)
    c.font = Font(name="Georgia", size=20, bold=True, color="FFFFFF")
    c.fill = fill(GREEN)
    c.alignment = Alignment(vertical="center", indent=1)
    for col in range(3, last_col + 1):
        ws.cell(2, col).fill = fill(GREEN)
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=last_col)
    s = ws.cell(3, 2, sub)
    s.font = font(italic=True, color=CRAN, size=10)
    s.alignment = Alignment(vertical="center", indent=1)

def header(ws, row, labels, start_col=2):
    for i, lab in enumerate(labels):
        c = ws.cell(row, start_col + i, lab)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = fill(CRAN)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
    ws.row_dimensions[row].height = 30

def cell(ws, r, c, v=None, inp=False, fmt=None, bold=False, color=None, align=None, wrap=False):
    x = ws.cell(r, c)
    if v is not None:
        x.value = v
    x.border = BOX
    x.fill = fill(INPUT if inp else "FFFFFF")
    x.font = font(bold=bold, color=color or ("1F3A93" if inp else "222222"))
    if fmt:
        x.number_format = fmt
    x.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    return x

def section(ws, r, text, c1=2, c2=6):
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    x = ws.cell(r, c1, text)
    x.font = font(bold=True, size=11, color=GREEN)
    x.border = Border(bottom=Side(style="medium", color=GOLD))
    for c in range(c1, c2 + 1):
        ws.cell(r, c).border = Border(bottom=Side(style="medium", color=GOLD))

def dv_list(ws, options, rng):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)

QTY = '#,##0.0;-#,##0.0;"–"'
INT = '#,##0;-#,##0;"–"'
MONEY = '#,##0.00;-#,##0.00;"–"'
TIME = "h:mm AM/PM"
DATE = "ddd, mmm d, yyyy"

wb.remove(wb.active)
start = new_sheet("Start Here", GREEN, {"B": 30, "C": 26, "D": 4, "E": 30, "F": 26})
setup = new_sheet("Setup", GOLD, {"B": 34, "C": 20, "D": 58})
guests = new_sheet("Guests", CRAN, {"B": 5, "C": 24, "D": 10, "E": 10, "F": 16, "G": 26, "H": 20, "I": 12})
food = new_sheet("Food", GREEN, {"B": 13, "C": 34, "D": 10, "E": 11, "F": 9, "G": 12, "H": 11, "I": 9, "J": 11, "K": 9, "L": 40})
drinks = new_sheet("Drinks", CRAN, {"B": 40, "C": 12, "D": 13, "E": 12, "F": 13, "G": 44})
shop = new_sheet("Shopping", GOLD, {"B": 11, "C": 34, "D": 10, "E": 13, "F": 12, "G": 12, "H": 10})
budget = new_sheet("Budget", GREEN, {"B": 26, "C": 14, "D": 14, "E": 14, "F": 30})
tl = new_sheet("Timeline", CRAN, {"B": 30, "C": 11, "D": 10, "E": 10, "F": 11, "G": 13, "H": 14, "I": 14, "J": 34})
chk = new_sheet("Checklist", GOLD, {"B": 16, "C": 60, "D": 10})

# ---------------------------------------------------------------- SETUP
title(setup, "Setup", "Fill in the gold cells. Everything else in the planner updates automatically.", 4)
S = {}
section(setup, 5, "Your event", 2, 4)
rows = [
    ("event", "Event name", "Thanksgiving Dinner", None, "Any name — shows on the Start Here page."),
    ("date", "Event date", dt.date(2026, 11, 26), DATE, "Used for the turkey-thawing and shopping reminders."),
    ("serve", "Dinner is served at", dt.time(17, 0), TIME, "The cooking timeline counts backwards from this time."),
    ("hours", "Party length (hours)", 5, "0", "Drives the drinks calculation."),
    ("source", "Guest count from", "Manual", None, "Manual = use the two numbers below. Guest List = count confirmed (RSVP Yes) guests on the Guests tab."),
    ("adults_m", "Adults (manual)", 10, "0", "Teens 12+ count as adults."),
    ("kids_m", "Kids under 12 (manual)", 2, "0", "Kids eat about half an adult portion (editable below)."),
    ("appetite", "Appetite of the crowd", "Normal", None, "Light = -15%, Normal, Hearty = +20% food."),
    ("leftovers", "Want leftovers?", "Yes", None, "Yes adds 25% to all food so everyone takes a plate home."),
    ("units", "Units for shopping", "Imperial (lb)", None, "Imperial (lb, qt) or Metric (kg, L). Both are always shown on the Food tab."),
    ("drinkers", "Share of adults who drink alcohol", 0.75, "0%", "Non-drinkers and kids get extra soft drinks instead."),
    ("budget", "Total budget", 600, "#,##0.00", "In your own currency. Tracked on the Budget tab."),
]
r = 6
for key, lab, val, fmt, note in rows:
    cell(setup, r, 2, lab, bold=True)
    cell(setup, r, 3, val, inp=True, fmt=fmt, align="center")
    n = setup.cell(r, 4, note)
    n.font = font(color=GREY, italic=True, size=9)
    n.alignment = Alignment(wrap_text=True, vertical="center")
    S[key] = f"Setup!$C${r}"
    setup.row_dimensions[r].height = 22
    r += 1
dv_list(setup, ["Manual", "Guest List"], S["source"].split("!")[1].replace("$", ""))
dv_list(setup, ["Light", "Normal", "Hearty"], S["appetite"].split("!")[1].replace("$", ""))
dv_list(setup, ["Yes", "No"], S["leftovers"].split("!")[1].replace("$", ""))
dv_list(setup, ["Imperial (lb)", "Metric (kg)"], S["units"].split("!")[1].replace("$", ""))

r += 1
section(setup, r, "Portion assumptions (change only if you want to)", 2, 4)
r += 1
for key, lab, val, fmt, note in [
    ("kidf", "Kid portion vs. adult", 0.5, "0%", "Common catering rule of thumb: a child under 12 eats about half an adult portion."),
    ("left_pct", "Extra food for leftovers", 0.25, "0%", "Applied only when 'Want leftovers?' is Yes."),
    ("light", "Light appetite adjustment", 0.85, "0%", "Multiplier on every portion."),
    ("hearty", "Hearty appetite adjustment", 1.2, "0%", "Multiplier on every portion."),
]:
    cell(setup, r, 2, lab, bold=True)
    cell(setup, r, 3, val, inp=True, fmt=fmt, align="center")
    n = setup.cell(r, 4, note)
    n.font = font(color=GREY, italic=True, size=9)
    n.alignment = Alignment(wrap_text=True, vertical="center")
    S[key] = f"Setup!$C${r}"
    r += 1

r += 1
section(setup, r, "Calculated (do not edit)", 2, 4)
r += 1
calc = [
    ("adults", "Adults used", f'=IF({S["source"]}="Guest List",Guests!$I$6,{S["adults_m"]})', "0"),
    ("kids", "Kids used", f'=IF({S["source"]}="Guest List",Guests!$I$7,{S["kids_m"]})', "0"),
    ("total", "Total guests", None, "0"),
    ("appmult", "Appetite multiplier", f'=IF({S["appetite"]}="Light",{S["light"]},IF({S["appetite"]}="Hearty",{S["hearty"]},1))', "0.00"),
    ("leftmult", "Leftover multiplier", f'=IF({S["leftovers"]}="Yes",1+{S["left_pct"]},1)', "0.00"),
    ("pf", "Portion factor (adult portions to cook)", None, "0.0"),
    ("metric", "Metric units? (1 = yes)", f'=IF(LEFT({S["units"]},6)="Metric",1,0)', "0"),
]
for key, lab, fml, fmt in calc:
    S[key] = f"Setup!$C${r}"
    if key == "total":
        fml = f'={S["adults"]}+{S["kids"]}'
    if key == "pf":
        fml = f'=({S["adults"]}+{S["kids"]}*{S["kidf"]})*{S["appmult"]}*{S["leftmult"]}'
    cell(setup, r, 2, lab)
    cell(setup, r, 3, fml, fmt=fmt, align="center", bold=True)
    r += 1
setup.freeze_panes = "A5"

# ---------------------------------------------------------------- GUESTS
title(guests, "Guest List", "Track RSVPs, diets and allergies. Set Setup → 'Guest count from' to Guest List to use these numbers.", 9)
header(guests, 9, ["#", "Name", "Adult / Kid", "RSVP", "Diet", "Allergies / notes", "Bringing", "Table / seat"])
for lab, rr in [("Confirmed adults", 6), ("Confirmed kids", 7)]:
    pass
cell(guests, 5, 8, "Summary", bold=True, color=GREEN)
summ = [
    (6, "Confirmed adults", '=COUNTIFS($D$10:$D$69,"Adult",$E$10:$E$69,"Yes")'),
    (7, "Confirmed kids", '=COUNTIFS($D$10:$D$69,"Kid",$E$10:$E$69,"Yes")'),
]
for rr, lab, fml in summ:
    cell(guests, rr, 8, lab)
    cell(guests, rr, 9, fml, fmt="0", align="center", bold=True)
cell(guests, 5, 3, "Waiting for reply", )
cell(guests, 5, 4, '=COUNTIF($E$10:$E$69,"Maybe")+COUNTIFS($C$10:$C$69,"<>",$E$10:$E$69,"")', fmt="0", align="center", bold=True)
cell(guests, 6, 3, "Special diets")
cell(guests, 6, 4, '=COUNTIFS($F$10:$F$69,"<>None",$F$10:$F$69,"<>",$E$10:$E$69,"Yes")', fmt="0", align="center", bold=True)
cell(guests, 7, 3, "Not coming")
cell(guests, 7, 4, '=COUNTIF($E$10:$E$69,"No")', fmt="0", align="center", bold=True)
examples = [
    ("Anna (host)", "Adult", "Yes", "None", "", "Turkey", "Head"),
    ("Mark", "Adult", "Yes", "Vegetarian", "Allergic to walnuts", "Wine", "T1"),
    ("Lily", "Kid", "Maybe", "None", "", "", "Kids' table"),
]
for i in range(60):
    rr = 10 + i
    cell(guests, rr, 2, i + 1, align="center", color=GREY)
    for c in range(3, 10):
        cell(guests, rr, c, inp=True, align="center" if c in (4, 5, 6, 9) else None)
    if i < len(examples):
        for j, v in enumerate(examples[i]):
            guests.cell(rr, 3 + j).value = v if v else None
dv_list(guests, ["Adult", "Kid"], "D10:D69")
dv_list(guests, ["Yes", "No", "Maybe"], "E10:E69")
dv_list(guests, ["None", "Vegetarian", "Vegan", "Gluten-free", "Dairy-free", "Pescatarian", "Other"], "F10:F69")
guests.conditional_formatting.add("E10:E69", CellIsRule(operator="equal", formula=['"Yes"'], fill=fill("D9EAD3")))
guests.conditional_formatting.add("E10:E69", CellIsRule(operator="equal", formula=['"No"'], fill=fill("F4CCCC")))
guests.freeze_panes = "A10"
guests.column_dimensions["I"].width = 14

# ---------------------------------------------------------------- FOOD
title(food, "Food Calculator", "Portions per ADULT for a sit-down holiday dinner with several dishes. Kids, appetite and leftovers are applied automatically.", 12)
cell(food, 5, 2, "Guests", bold=True)
cell(food, 5, 3, f'={S["adults"]}&" adults + "&{S["kids"]}&" kids"', bold=True, color=GREEN)
cell(food, 6, 2, "Cooking for", bold=True)
cell(food, 6, 3, f'=TEXT({S["pf"]},"0.0")&" adult portions"', bold=True, color=GREEN)
header(food, 8, ["Course", "Item", "Include?", "Per adult", "Unit", "Total needed", "Metric", "", "Imperial", "", "Tip"])
food.merge_cells("H8:I8")
food.merge_cells("J8:K8")
items = [
    ("Main", "Whole turkey (bone-in, raw weight)", "Yes", 550, "g", "About 1¼ lb per person bone-in. Cook to 165°F / 74°C."),
    ("Main", "Ham (bone-in)", "No", 250, "g", "Use 150 g if also serving turkey."),
    ("Main", "Roast beef / prime rib (boneless)", "No", 200, "g", "Raw weight; shrinks ~25% when cooked."),
    ("Main", "Whole chicken (raw)", "No", 450, "g", ""),
    ("Main", "Salmon / fish fillet", "No", 170, "g", ""),
    ("Main", "Vegetarian main (e.g. nut roast)", "No", 200, "g", "Tip: set per adult so it covers only veggie guests."),
    ("Side", "Potatoes (mashed or roasted, raw)", "Yes", 200, "g", "Most popular side — people take seconds."),
    ("Side", "Stuffing / dressing (bread cubes)", "Yes", 60, "g", "Dry bread weight; ~1 cup prepared per person."),
    ("Side", "Vegetable side #1 (green beans)", "Yes", 120, "g", ""),
    ("Side", "Vegetable side #2 (Brussels sprouts / carrots)", "Yes", 120, "g", ""),
    ("Side", "Sweet potatoes", "No", 150, "g", ""),
    ("Side", "Mac & cheese (dry pasta)", "No", 50, "g", ""),
    ("Side", "Salad greens", "No", 40, "g", ""),
    ("Side", "Gravy", "Yes", 90, "ml", "About ⅓ cup per person."),
    ("Side", "Cranberry sauce", "Yes", 60, "g", ""),
    ("Side", "Dinner rolls", "Yes", 1.5, "rolls", ""),
    ("Starter", "Cheese (for board)", "Yes", 50, "g", "Use 100 g if the board IS the meal."),
    ("Starter", "Cured meats", "No", 40, "g", ""),
    ("Starter", "Crackers / baguette", "Yes", 40, "g", ""),
    ("Starter", "Dips + veggie sticks", "Yes", 80, "g", ""),
    ("Starter", "Nuts / olives / snacks", "No", 30, "g", ""),
    ("Dessert", "Pies (8 slices each)", "Yes", 0.16, "pies", "1¼ slices per person — someone always wants a second piece."),
    ("Dessert", "Cake (12 slices)", "No", 0.1, "cakes", ""),
    ("Dessert", "Whipped cream / ice cream", "Yes", 60, "ml", ""),
    ("Dessert", "Coffee (ground)", "Yes", 10, "g", "About 1 cup per adult."),
    ("Custom", "", "No", None, "g", "Add your own dish: name, per adult, unit (g / ml / any word for pieces)."),
    ("Custom", "", "No", None, "g", ""),
    ("Custom", "", "No", None, "g", ""),
    ("Custom", "", "No", None, "g", ""),
]
FOOD_FIRST = 9
course_fill = {"Main": "F4E3E6", "Side": "E7F0EA", "Starter": "FBF0D9", "Dessert": "EFE7F4", "Custom": "FFFFFF"}
for i, (course, name, inc, per, unit, tip) in enumerate(items):
    rr = FOOD_FIRST + i
    c = cell(food, rr, 2, course, align="center", bold=True, color=GREEN)
    c.fill = fill(course_fill[course])
    cell(food, rr, 3, name or None, inp=(course == "Custom"))
    cell(food, rr, 4, inc, inp=True, align="center")
    cell(food, rr, 5, per, inp=True, align="center", fmt="General")
    cell(food, rr, 6, unit, inp=(course == "Custom"), align="center")
    cell(food, rr, 7, f'=IF(OR(D{rr}<>"Yes",E{rr}=""),0,IF(OR(F{rr}="g",F{rr}="ml"),E{rr}*{S["pf"]},ROUNDUP(E{rr}*{S["pf"]},0)))', fmt=INT, align="center")
    cell(food, rr, 8, f'=IF(OR(F{rr}="g",F{rr}="ml"),G{rr}/1000,G{rr})', fmt=QTY, align="center", bold=True)
    cell(food, rr, 9, f'=IF(F{rr}="g","kg",IF(F{rr}="ml","L",F{rr}))', align="center", color=GREY)
    cell(food, rr, 10, f'=IF(F{rr}="g",G{rr}/453.6,IF(F{rr}="ml",G{rr}/946.4,G{rr}))', fmt=QTY, align="center", bold=True)
    cell(food, rr, 11, f'=IF(F{rr}="g","lb",IF(F{rr}="ml","qt",F{rr}))', align="center", color=GREY)
    t = cell(food, rr, 12, tip or None)
    t.font = font(color=GREY, italic=True, size=9)
FOOD_LAST = FOOD_FIRST + len(items) - 1
FR = f"{FOOD_FIRST}:{FOOD_LAST}"
dv_list(food, ["Yes", "No"], f"D{FOOD_FIRST}:D{FOOD_LAST}")
food.conditional_formatting.add(f"C{FOOD_FIRST}:K{FOOD_LAST}",
    FormulaRule(formula=[f'$D{FOOD_FIRST}<>"Yes"'], font=Font(name=F, color="B7B7B7")))
tr = FOOD_LAST + 2
cell(food, tr, 3, "Total food to buy (by weight)", bold=True)
cell(food, tr, 8, f'=SUMIF(F{FOOD_FIRST}:F{FOOD_LAST},"g",G{FOOD_FIRST}:G{FOOD_LAST})/1000', fmt=QTY, bold=True, align="center")
cell(food, tr, 9, "kg", align="center", color=GREY)
cell(food, tr, 10, f"=H{tr}*2.2046", fmt=QTY, bold=True, align="center")
cell(food, tr, 11, "lb", align="center", color=GREY)
FOOD_TOTAL_KG = f"Food!$H${tr}"
TURKEY_LB = f"Food!$J${FOOD_FIRST}"
n = food.cell(tr + 2, 3, "Portion sizes are standard catering rules of thumb for a multi-dish holiday meal. Adjust 'Per adult' to match your crowd.")
n.font = font(italic=True, color=GREY, size=9)
food.freeze_panes = f"A{FOOD_FIRST}"

# ---------------------------------------------------------------- DRINKS
title(drinks, "Drinks Calculator", "Based on the classic host rule: 2 drinks in the first hour, then 1 drink per hour per guest.", 7)
section(drinks, 5, "Assumptions (gold = editable)", 2, 7)
D = {}
dr = [
    ("first", "Drinks per drinker — first hour", 2, "0", "The classic party rule of thumb."),
    ("next", "Drinks per drinker — each extra hour", 1, "0", ""),
    ("wine", "Share of drinks that are wine", 0.5, "0%", "Wine + beer + cocktails should add up to 100%."),
    ("beer", "Share that is beer / cider", 0.3, "0%", ""),
    ("spirit", "Share that is cocktails / spirits", 0.2, "0%", ""),
    ("glasses", "Glasses per bottle of wine (750 ml)", 5, "0", "5 glasses of 150 ml / 5 oz."),
    ("shots", "Drinks per bottle of spirits (750 ml)", 16, "0", "About 1.5 oz / 45 ml per drink."),
    ("mixer", "Mixer per cocktail (L)", 0.15, "0.00", "About 5 oz."),
    ("soft", "Soft drink serving (L)", 0.25, "0.00", "About 8 oz glass."),
    ("water", "Water per guest (L)", 0.5, "0.0", "Still + sparkling for the table."),
    ("ice", "Ice per guest (kg)", 0.5, "0.0", "About 1 lb per guest. Double it on a hot day."),
]
rr = 6
for key, lab, val, fmt, note in dr:
    cell(drinks, rr, 2, lab)
    cell(drinks, rr, 3, val, inp=True, fmt=fmt, align="center")
    drinks.merge_cells(start_row=rr, start_column=4, end_row=rr, end_column=7)
    nn = drinks.cell(rr, 4, note)
    nn.font = font(italic=True, color=GREY, size=9)
    D[key] = f"$C${rr}"
    rr += 1
cell(drinks, rr, 2, "Check: shares add up to", bold=True)
cell(drinks, rr, 3, f'={D["wine"]}+{D["beer"]}+{D["spirit"]}', fmt="0%", align="center", bold=True)
drinks.conditional_formatting.add(f"C{rr}", CellIsRule(operator="notEqual", formula=["1"], fill=fill("F4CCCC"), font=Font(name=F, bold=True, color=CRAN)))
rr += 2
section(drinks, rr, "What to buy", 2, 7)
rr += 1
header(drinks, rr, ["Item", "Metric", "", "Imperial", "", "How it's calculated"])
drinks.merge_cells(start_row=rr, start_column=3, end_row=rr, end_column=4)
drinks.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=6)
rr += 1
A, K, H, DK = S["adults"], S["kids"], S["hours"], S["drinkers"]
alc = f'({A}*{DK}*({D["first"]}+MAX({H}-1,0)*{D["next"]}))'
DR_FIRST = rr
drink_rows = [
    ("Alcoholic drinks in total", f"=ROUND({alc},0)", "drinks", f"=C{{r}}", "drinks", "Adults who drink × drinks per hour × hours."),
    ("Wine", f'=ROUNDUP({alc}*{D["wine"]}/{D["glasses"]},0)', "bottles", "=C{r}", "bottles", "750 ml bottles."),
    ("Beer / cider", f'=ROUNDUP({alc}*{D["beer"]},0)', "bottles", "=C{r}", "bottles", "330 ml / 12 oz bottles or cans."),
    ("Spirits for cocktails", f'=ROUNDUP({alc}*{D["spirit"]}/{D["shots"]},0)', "bottles", "=C{r}", "bottles", "750 ml bottles."),
    ("Mixers (tonic, soda, juice)", f'=ROUNDUP({alc}*{D["spirit"]}*{D["mixer"]},1)', "L", "=C{r}*1.0567", "qt", ""),
    ("Soft drinks & juice", f'=ROUNDUP(({A}*(1-{DK})+{K})*({D["first"]}+MAX({H}-1,0)*{D["next"]})*{D["soft"]}+{A}*{DK}*{D["soft"]},1)', "L", "=C{r}*1.0567", "qt", "Kids + non-drinkers drink at the same pace; drinkers get one glass too."),
    ("  → in 2-liter bottles", "=ROUNDUP(C{p}/2,0)", "bottles", "=C{r}", "bottles", ""),
    ("Water", f'=ROUNDUP(({A}+{K})*{D["water"]},1)', "L", "=C{r}*1.0567", "qt", ""),
    ("Ice", f'=ROUNDUP(({A}+{K})*{D["ice"]},1)', "kg", "=C{r}*2.2046", "lb", ""),
]
for i, (lab, m, mu, imp, iu, note) in enumerate(drink_rows):
    rrr = rr + i
    cell(drinks, rrr, 2, lab, bold=(i == 0))
    cell(drinks, rrr, 3, m.replace("{p}", str(rrr - 1)), fmt=QTY if mu in ("L", "kg") else INT, align="center", bold=True)
    cell(drinks, rrr, 4, mu, align="center", color=GREY)
    cell(drinks, rrr, 5, imp.replace("{r}", str(rrr)), fmt=QTY if iu in ("qt", "lb") else INT, align="center", bold=True)
    cell(drinks, rrr, 6, iu, align="center", color=GREY)
    nn = cell(drinks, rrr, 7, note or None)
    nn.font = font(italic=True, color=GREY, size=9)
DRINK_ROWS = {lab.strip(" →"): rr + i for i, (lab, *_rest) in enumerate(drink_rows)}

# ---------------------------------------------------------------- SHOPPING
title(shop, "Shopping List", "Quantities fill in automatically. Add your store's prices to see the cost. Tick items as you buy them.", 8)
cell(shop, 5, 2, "Units", bold=True)
cell(shop, 5, 3, f'={S["units"]}', bold=True, color=GREEN)
cell(shop, 6, 2, "Est. total", bold=True)
header(shop, 8, ["Category", "Item", "Qty", "Unit", "Price / unit", "Est. cost", "Bought?"])
sr = 9
SHOP_FIRST = sr
ex_price = {  # example prices, US grocery, per imperial unit
    0: 1.99, 6: 0.99, 7: 3.50, 8: 2.49, 9: 3.49, 13: 2.00, 14: 3.99, 15: 0.50,
    16: 9.99, 18: 4.99, 19: 4.99, 21: 12.99, 23: 5.00, 24: 8.99,
}
for i in range(len(items)):
    fr = FOOD_FIRST + i
    cell(shop, sr, 2, "Food", align="center", color=GREEN)
    cell(shop, sr, 3, f'=IF(Food!C{fr}="","",Food!C{fr})')
    cell(shop, sr, 4, f'=IF({S["metric"]}=1,Food!H{fr},Food!J{fr})', fmt=QTY, align="center", bold=True)
    cell(shop, sr, 5, f'=IF({S["metric"]}=1,Food!I{fr},Food!K{fr})', align="center", color=GREY)
    cell(shop, sr, 6, ex_price.get(i), inp=True, fmt=MONEY, align="center")
    cell(shop, sr, 7, f"=D{sr}*F{sr}", fmt=MONEY, align="center")
    cell(shop, sr, 8, None, inp=True, align="center")
    sr += 1
drink_price = {"Wine": 12.00, "Beer / cider": 1.50, "Spirits for cocktails": 22.00, "in 2-liter bottles": 2.29, "Ice": 0.60}
for lab, (dlab, *_r) in zip(DRINK_ROWS.keys(), drink_rows):
    if lab in ("Alcoholic drinks in total", "Soft drinks & juice"):
        continue
    d_r = DRINK_ROWS[lab]
    name = "Soft drinks & juice (2-liter bottles)" if lab == "in 2-liter bottles" else lab
    cell(shop, sr, 2, "Drinks", align="center", color=CRAN)
    cell(shop, sr, 3, name)
    cell(shop, sr, 4, f"=IF({S['metric']}=1,Drinks!C{d_r},Drinks!E{d_r})", fmt=QTY, align="center", bold=True)
    cell(shop, sr, 5, f"=IF({S['metric']}=1,Drinks!D{d_r},Drinks!F{d_r})", align="center", color=GREY)
    cell(shop, sr, 6, drink_price.get(lab), inp=True, fmt=MONEY, align="center")
    cell(shop, sr, 7, f"=D{sr}*F{sr}", fmt=MONEY, align="center")
    cell(shop, sr, 8, None, inp=True, align="center")
    sr += 1
extras = [("Other", "Butter", 2, "lb", 4.50), ("Other", "Heavy cream", 1, "qt", 5.99), ("Other", "Napkins", 1, "pack", 4.00)]
for i in range(12):
    cat, name, q, u, p = extras[i] if i < len(extras) else ("Other", None, None, None, None)
    cell(shop, sr, 2, cat, inp=True, align="center")
    cell(shop, sr, 3, name, inp=True)
    cell(shop, sr, 4, q, inp=True, fmt=QTY, align="center")
    cell(shop, sr, 5, u, inp=True, align="center")
    cell(shop, sr, 6, p, inp=True, fmt=MONEY, align="center")
    cell(shop, sr, 7, f"=IF(D{sr}=\"\",0,D{sr}*F{sr})", fmt=MONEY, align="center")
    cell(shop, sr, 8, None, inp=True, align="center")
    sr += 1
SHOP_LAST = sr - 1
dv_list(shop, ["✓"], f"H{SHOP_FIRST}:H{SHOP_LAST}")
dv_list(shop, ["Food", "Drinks", "Other", "Decor", "Tableware"], f"B{SHOP_LAST-11}:B{SHOP_LAST}")
shop.conditional_formatting.add(f"B{SHOP_FIRST}:H{SHOP_LAST}",
    FormulaRule(formula=[f'$H{SHOP_FIRST}="✓"'], fill=fill("D9EAD3"), font=Font(name=F, color="6AA84F", strike=True)))
shop.conditional_formatting.add(f"B{SHOP_FIRST}:H{SHOP_LAST}",
    FormulaRule(formula=[f'AND($C{SHOP_FIRST}<>"",$D{SHOP_FIRST}=0)'], font=Font(name=F, color="C8C8C8")))
cell(shop, 6, 3, f"=SUM(G{SHOP_FIRST}:G{SHOP_LAST})", fmt="#,##0.00", bold=True, color=CRAN)
cell(shop, 5, 5, "Still to buy", bold=True)
cell(shop, 5, 6, f'=COUNTIFS(H{SHOP_FIRST}:H{SHOP_LAST},"<>✓",D{SHOP_FIRST}:D{SHOP_LAST},">0")', fmt="0", align="center", bold=True, color=CRAN)
cell(shop, 6, 5, "Bought", bold=True)
cell(shop, 6, 6, f'=COUNTIF(H{SHOP_FIRST}:H{SHOP_LAST},"✓")', fmt="0", align="center", bold=True, color=GREEN)
note = shop.cell(SHOP_LAST + 2, 3, "Example prices are typical US grocery prices (2026) per unit shown — replace them with your own store's prices. Grey rows = not needed.")
note.font = font(italic=True, color=GREY, size=9)
shop.freeze_panes = f"A{SHOP_FIRST}"
SHOP_TOTAL = "Shopping!$C$6"

# ---------------------------------------------------------------- BUDGET
title(budget, "Budget", "Planned costs come from the Shopping List. Type what you actually spent in the gold column.", 6)
header(budget, 5, ["Category", "Planned", "Actual", "Difference", "Notes"])
cats = [
    ("Food", f'=SUMIF(Shopping!$B${SHOP_FIRST}:$B${SHOP_LAST},"Food",Shopping!$G${SHOP_FIRST}:$G${SHOP_LAST})', "From Shopping List"),
    ("Drinks", f'=SUMIF(Shopping!$B${SHOP_FIRST}:$B${SHOP_LAST},"Drinks",Shopping!$G${SHOP_FIRST}:$G${SHOP_LAST})', "From Shopping List"),
    ("Other groceries", f'=SUMIF(Shopping!$B${SHOP_FIRST}:$B${SHOP_LAST},"Other",Shopping!$G${SHOP_FIRST}:$G${SHOP_LAST})', "From Shopping List"),
    ("Decor & flowers", 40, ""),
    ("Tableware & linens", 30, ""),
    ("Gifts / party favors", 0, ""),
    ("Miscellaneous", 25, ""),
]
br = 6
for lab, planned, nte in cats:
    cell(budget, br, 2, lab, bold=True)
    auto = isinstance(planned, str)
    cell(budget, br, 3, planned, inp=not auto, fmt=MONEY, align="center")
    cell(budget, br, 4, None, inp=True, fmt=MONEY, align="center")
    cell(budget, br, 5, f'=IF(D{br}="","",C{br}-D{br})', fmt=MONEY, align="center")
    nn = cell(budget, br, 6, nte or None)
    nn.font = font(italic=True, color=GREY, size=9)
    br += 1
BL = br - 1
cell(budget, br, 2, "TOTAL", bold=True, color=GREEN)
cell(budget, br, 3, f"=SUM(C6:C{BL})", fmt=MONEY, align="center", bold=True)
cell(budget, br, 4, f"=SUM(D6:D{BL})", fmt=MONEY, align="center", bold=True)
cell(budget, br, 5, f'=IF(D{br}=0,"",C{br}-D{br})', fmt=MONEY, align="center", bold=True)
TOT = br
br += 2
section(budget, br, "Summary", 2, 6)
br += 1
summary = [
    ("Your budget", f'={S["budget"]}'),
    ("Planned spend", f"=C{TOT}"),
    ("Left over (planned)", f"=C{br}-C{br+1}"),
    ("Cost per guest (planned)", f'=IF({S["total"]}=0,0,C{br+1}/{S["total"]})'),
    ("Actual spend so far", f"=D{TOT}"),
]
for i, (lab, fml) in enumerate(summary):
    cell(budget, br + i, 2, lab, bold=True)
    cell(budget, br + i, 3, fml, fmt=MONEY, align="center", bold=True)
LEFT = f"Budget!$C${br+2}"
budget.conditional_formatting.add(f"C{br+2}", CellIsRule(operator="lessThan", formula=["0"], fill=fill("F4CCCC"), font=Font(name=F, bold=True, color=CRAN)))
budget.conditional_formatting.add(f"C{br+2}", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=fill("D9EAD3"), font=Font(name=F, bold=True, color=GREEN)))
budget.conditional_formatting.add(f"E6:E{TOT}", CellIsRule(operator="lessThan", formula=["0"], font=Font(name=F, bold=True, color=CRAN)))

# ---------------------------------------------------------------- TIMELINE
title(tl, "Cooking Timeline", "Counts backwards from the time dinner is served. Turkey time is calculated from its weight.", 10)
cell(tl, 5, 2, "Dinner is served at", bold=True)
cell(tl, 5, 3, f'={S["serve"]}', fmt=TIME, bold=True, color=CRAN, align="center")
cell(tl, 6, 2, "Turkey weight (lb)", bold=True)
cell(tl, 6, 3, f"={TURKEY_LB}", fmt=QTY, bold=True, align="center")
header(tl, 8, ["Dish", "Where", "Prep (min)", "Cook (min)", "Rest / hold (min)", "Start prep", "Start cooking", "Done cooking", "Notes"])
tasks = [
    ("Turkey", "Oven", 30, f"=ROUND({TURKEY_LB}*15,0)", 40, "~15 min per lb at 325°F / 165°C (unstuffed). Check 165°F / 74°C in the thigh."),
    ("Stuffing / dressing", "Oven", 20, 45, 10, "Bake while the turkey rests."),
    ("Sweet potatoes / casserole", "Oven", 20, 35, 5, ""),
    ("Roasted vegetables", "Oven", 15, 30, 5, ""),
    ("Mashed potatoes", "Stove", 20, 25, 0, "Keep warm in a slow cooker."),
    ("Green beans", "Stove", 10, 12, 0, ""),
    ("Gravy", "Stove", 5, 15, 0, "Use the turkey drippings."),
    ("Dinner rolls", "Oven", 5, 12, 0, "Warm up when the turkey comes out."),
    ("Appetizer board", "Counter", 25, 0, 90, "Ready when guests arrive (90 min before dinner)."),
    ("Pies", "Oven", 30, 55, 240, "Bake in the morning or the day before."),
    ("", "", None, None, None, ""),
    ("", "", None, None, None, ""),
    ("", "", None, None, None, ""),
]
TL_FIRST = 9
for i, (dish, where, prep, cook, rest, nte) in enumerate(tasks):
    t = TL_FIRST + i
    cell(tl, t, 2, dish or None, inp=True)
    cell(tl, t, 3, where or None, inp=True, align="center")
    cell(tl, t, 4, prep, inp=True, align="center")
    auto = isinstance(cook, str)
    cell(tl, t, 5, cook, inp=not auto, align="center", fmt="0")
    cell(tl, t, 6, rest, inp=True, align="center")
    serve = f"$C$5"
    cell(tl, t, 7, f'=IF(B{t}="","",MOD({serve}-(D{t}+E{t}+F{t})/1440,1))', fmt=TIME, align="center", bold=True, color=GREEN)
    cell(tl, t, 8, f'=IF(OR(B{t}="",E{t}=0),"",MOD({serve}-(E{t}+F{t})/1440,1))', fmt=TIME, align="center", bold=True, color=CRAN)
    cell(tl, t, 9, f'=IF(OR(B{t}="",E{t}=0),"",MOD({serve}-F{t}/1440,1))', fmt=TIME, align="center")
    nn = cell(tl, t, 10, nte or None, wrap=True)
    nn.font = font(italic=True, color=GREY, size=9)
    tl.row_dimensions[t].height = 26
TL_LAST = TL_FIRST + len(tasks) - 1
dv_list(tl, ["Oven", "Stove", "Grill", "Slow cooker", "Microwave", "Fridge", "Counter"], f"C{TL_FIRST}:C{TL_LAST}")
kr = TL_LAST + 2
section(tl, kr, "Key dates", 2, 10)
kr += 1
dates = [
    ("Move turkey to the fridge to thaw", f'=IF({TURKEY_LB}=0,"No turkey",{S["date"]}-ROUNDUP({TURKEY_LB}/4,0))', "1 day in the fridge for every 4 lb (USDA guideline)."),
    ("Big shopping trip", f'={S["date"]}-3', "Buy fresh produce 1–2 days before."),
    ("Bake pies, make cranberry sauce", f'={S["date"]}-1', ""),
    ("Party day", f'={S["date"]}', ""),
]
for i, (lab, fml, nte) in enumerate(dates):
    cell(tl, kr + i, 2, lab, bold=True)
    tl.merge_cells(start_row=kr + i, start_column=3, end_row=kr + i, end_column=5)
    cell(tl, kr + i, 3, fml, fmt=DATE, bold=True, color=CRAN, align="center")
    for c in (4, 5):
        tl.cell(kr + i, c).border = BOX
    tl.merge_cells(start_row=kr + i, start_column=6, end_row=kr + i, end_column=10)
    nn = tl.cell(kr + i, 6, nte)
    nn.font = font(italic=True, color=GREY, size=9)
THAW = f"Timeline!$C${kr}"
TURKEY_START = f"Timeline!$H${TL_FIRST}"
tl.freeze_panes = f"A{TL_FIRST}"

# ---------------------------------------------------------------- CHECKLIST
title(chk, "Host Checklist", "Pick a ✓ in the Done column as you go.", 4)
header(chk, 5, ["When", "Task", "Done"])
tasks = {
    "3 weeks before": ["Set the date and time, send invitations", "Ask about diets and allergies (Guests tab)", "Decide the menu (Food tab: Include? = Yes)", "Order the turkey if buying fresh / heritage"],
    "2 weeks before": ["Check you have enough chairs, plates, glasses, cutlery", "Plan the cooking timeline — does everything fit in the oven?", "Buy drinks and non-perishables"],
    "1 week before": ["Confirm RSVPs and update guest count", "Plan seating and table layout", "Clean out the fridge to make room"],
    "3–4 days before": ["Start thawing a frozen turkey (see Key dates)", "Big shopping trip (Shopping tab)", "Iron tablecloth and napkins"],
    "Day before": ["Bake pies and make cranberry sauce", "Chop vegetables, make stuffing base", "Set the table", "Chill drinks, buy ice"],
    "Party day": ["Follow the Cooking Timeline", "Set up appetizers and drinks station", "Light candles, put on music", "Enjoy your guests!"],
    "After": ["Pack leftovers for guests", "Note what ran out / was left for next year"],
}
cr = 6
for when, lst in tasks.items():
    for j, tsk in enumerate(lst):
        c = cell(chk, cr, 2, when if j == 0 else None, bold=True, color=GREEN)
        c.fill = fill("F3EDE0")
        cell(chk, cr, 3, tsk)
        cell(chk, cr, 4, None, inp=True, align="center")
        cr += 1
for i in range(5):
    cell(chk, cr, 2, None, inp=True)
    cell(chk, cr, 3, None, inp=True)
    cell(chk, cr, 4, None, inp=True, align="center")
    cr += 1
dv_list(chk, ["✓"], f"D6:D{cr-1}")
chk.conditional_formatting.add(f"C6:D{cr-1}", FormulaRule(formula=['$D6="✓"'], fill=fill("D9EAD3"), font=Font(name=F, color="6AA84F", strike=True)))
cell(chk, 4, 3, f'=COUNTIF(D6:D{cr-1},"✓")&" of "&COUNTA(C6:C{cr-1})&" tasks done"', bold=True, color=CRAN)
chk.freeze_panes = "A6"

# ---------------------------------------------------------------- START HERE
st = start
title(st, "Holiday Hosting Planner", "Guest count in → food, drinks, shopping list, budget and cooking schedule out.", 6)
section(st, 5, "Your party at a glance", 2, 6)
glance = [
    ("Event", f'={S["event"]}', None),
    ("Date", f'={S["date"]}', DATE),
    ("Dinner is served", f'={S["serve"]}', TIME),
    ("Guests", f'={S["adults"]}&" adults, "&{S["kids"]}&" kids"', None),
    ("Food to buy", f'=TEXT({FOOD_TOTAL_KG},"0.0")&" kg  /  "&TEXT({FOOD_TOTAL_KG}*2.2046,"0.0")&" lb"', None),
]
glance2 = [
    ("Estimated shopping", f"={SHOP_TOTAL}", "#,##0.00"),
    ("Budget left", f"={LEFT}", "#,##0.00"),
    ("Turkey goes in the oven", f'=IF({TURKEY_LB}=0,"—",{TURKEY_START})', TIME),
    ("Start thawing turkey", f"={THAW}", DATE),
    ("Checklist", f"=Checklist!C4", None),
]
for i, (lab, fml, fmt) in enumerate(glance):
    cell(st, 6 + i, 2, lab, bold=True)
    cell(st, 6 + i, 3, fml, fmt=fmt, bold=True, color=CRAN, align="center")
for i, (lab, fml, fmt) in enumerate(glance2):
    cell(st, 6 + i, 5, lab, bold=True)
    cell(st, 6 + i, 6, fml, fmt=fmt, bold=True, color=CRAN, align="center")
for rrr in range(6, 11):
    st.row_dimensions[rrr].height = 22

section(st, 13, "How to use it (3 steps)", 2, 6)
steps = [
    "1.  Setup tab — type your date, dinner time and number of guests (gold cells).",
    "2.  Food tab — set Include? to Yes for the dishes on your menu. Quantities appear instantly.",
    "3.  Shopping tab — add your store's prices, then tick items off as you buy them. Budget, Timeline and Checklist update by themselves.",
]
for i, s_ in enumerate(steps):
    st.merge_cells(start_row=14 + i, start_column=2, end_row=14 + i, end_column=6)
    x = st.cell(14 + i, 2, s_)
    x.font = font(size=10)
    x.alignment = Alignment(wrap_text=True, vertical="center")
    st.row_dimensions[14 + i].height = 22

section(st, 18, "Colour guide", 2, 6)
cell(st, 19, 2, "Gold cell", inp=True)
st.cell(19, 3, "Type or pick here").font = font(color=GREY)
cell(st, 20, 2, "White cell")
st.cell(20, 3, "Calculated for you — no need to touch").font = font(color=GREY)

section(st, 22, "What's inside", 2, 6)
inside = [
    ("Setup", "Date, time, guests, appetite, leftovers, units, budget"),
    ("Guests", "RSVPs, diets, allergies, who brings what, seating"),
    ("Food", "25 holiday dishes + your own, per-person portions → exact kg / lb"),
    ("Drinks", "Wine, beer, spirits, soft drinks, water, ice"),
    ("Shopping", "Auto-filled list with prices and a bought ✓"),
    ("Budget", "Planned vs. actual, cost per guest"),
    ("Timeline", "When to start each dish + turkey thaw date"),
    ("Checklist", "3 weeks out → party day, tick as you go"),
]
for i, (a, b) in enumerate(inside):
    cell(st, 23 + i, 2, a, bold=True, color=GREEN)
    st.merge_cells(start_row=23 + i, start_column=3, end_row=23 + i, end_column=6)
    cell(st, 23 + i, 3, b)
    for c in range(4, 7):
        st.cell(23 + i, c).border = BOX
tip = st.cell(32, 2, "Works in Microsoft Excel and Google Sheets (File → Import, or upload to Google Drive and open with Google Sheets). Portions are rules of thumb — trust your crowd.")
st.merge_cells("B32:F33")
tip.font = font(italic=True, color=GREY, size=9)
tip.alignment = Alignment(wrap_text=True, vertical="top")

wb.active = 0
wb.save(OUT)
print("saved", OUT)
